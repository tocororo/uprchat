from pathlib import Path
from typing import List
from bs4 import BeautifulSoup
import fitz
from langchain_community.document_loaders import (
    PyMuPDFLoader,
)
from langchain_core.documents import Document as LangchainDocument
from docx import Document
import openai
from pptx import Presentation
from langchain_openai import ChatOpenAI
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain.chains.summarize import load_summarize_chain
from langchain.prompts import PromptTemplate

from uprchat.harvester.logger import setup_logger
from uprchat.utils import apikey_iterator

logger = setup_logger(__name__)

apikey_iterator = apikey_iterator.APIKeyIterator()

class Parser:
    """
    Parser for extracting text summaries from various document types.
    """

    def parse_pdf(self, path: Path) -> str:
        """
        Extracts text from the first page of a PDF file.
        Args:
            path:(Path)  to the PDF file.
        Returns:
            str: Extracted text summary.
        """
        try:
            with fitz.open(path) as doc:
                first_page = doc.load_page(0)
                return first_page.get_text()
        except Exception as e:
            logger.error(f"Error parsing PDF {path}: {e}")
            return ""

    def parse_docx(self, path: Path) -> str:
        """
        Extracts text from the first 100 paragraphs of a DOCX file.
        Args:
            path:(Path)  to the DOCX file.
        Returns:
            str: Extracted text summary.
        """
        try:
            with open(path, "rb") as f:
                document = Document(f)
                summary = " ".join(p.text for p in document.paragraphs[:100])
                return summary
        except Exception as e:
            logger.error(f"Error parsing DOCX {path}: {e}")
            return ""

    def parse_pptx(self, path: Path) -> str:
        """
        Extracts text from the first 100 slides of a PPTX file.
        Args:
            path:(Path)  to the PPTX file.
        Returns:
            str: Extracted text summary.
        """
        try:
            with open(path, "rb") as f:
                presentation = Presentation(f)
                summary = " ".join(
                    [f"{shape.text}" 
                     for slide in list(presentation.slides)[:100] 
                     for shape in slide.shapes
                     if shape.has_text_frame
                     ]
                    )
                return summary
        except Exception as e:
            logger.error(f"Error parsing PPTX {path}: {e}")
            return ""

class ParserAI:
    """
    Parser for extracting text summaries from various document types using LLMs.
    """
    def __init__(
        self,
        model_name: str = None,
        model_type: str = "openai",
        base_url: str = None,
        api_key: str = None
    ):
        if model_type != "openai":
            raise ValueError(f"Model type '{model_type}' not supported yet.")

        if not model_name:
            self.agent = apikey_iterator.get_llm()
        else:
            self.agent = ChatOpenAI(
                model_name=model_name, base_url=base_url, api_key=api_key
            )

        prompt = """
                Elabora un resumen en español del siguiente texto. Debe ser breve (alrededor de un
                párrafo), pero suficientemente detallado para captar las ideas principales:
                {text}
                """
        self.prompt_template = PromptTemplate(
            template=prompt, input_variables=["text"]
        )

    def _summarize(self, documents: List[LangchainDocument]) -> str:
        splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200)
        chunks = splitter.split_documents(documents)
        chain = load_summarize_chain(
            self.agent,
            chain_type="map_reduce",
            map_prompt=self.prompt_template,
            combine_prompt=self.prompt_template,
        )
        return chain.invoke(chunks)["output_text"]

    def parse_pdf(self, path: Path) -> str:
        try:
            docs = PyMuPDFLoader(path).load()
            return self._summarize(docs)
        except openai.RateLimitError as e:
            raise e
        except Exception as e:
            logger.error(f"Error parsing PDF {path}: {e}")
            return ""

    def parse_docx(self, path: Path) -> str:
        try:
            doc = Document(str(path))
            texts = []
            for paragraph in doc.paragraphs:
                text = paragraph.text
                texts.append(text)
            docs = [LangchainDocument(page_content=text) for text in texts]
            return self._summarize(docs)
        except openai.RateLimitError as e:
            raise e
        except Exception as e:
            logger.error(f"Error parsing DOCX {path}: {e}")
            return ""

    def parse_pptx(self, path: Path) -> str:
        try:
            prs = Presentation(str(path))
            texts = []
            for slide in prs.slides:
                slide_text = " ".join(
                    [shape.text for shape in slide.shapes if shape.has_text_frame ]
                )
                texts.append(slide_text)

            docs = [LangchainDocument(page_content=text) for text in texts]
            return self._summarize(docs)
        except openai.RateLimitError as e:
            raise e
        except Exception as e:
            logger.error(f"Error parsing PPTX {path}: {e}")
            return ""
        

    def parse_html(self, path: Path) -> str:
        """
        Parse and summarize the textual content of an HTML document.
        """
        try:
            with open(path, "r", encoding="utf-8") as file:
                html_content = file.read()

            soup = BeautifulSoup(html_content, "html.parser")

            for element in soup(["script", "style", "head", "meta", "noscript"]):
                element.decompose()

            text = soup.get_text(separator=" ", strip=True)
            doc = LangchainDocument(page_content=text)
            return self._summarize([doc])
        except openai.RateLimitError as e:
            raise e
        except Exception as e:
            logger.error(f"Error parsing HTML {path}: {e}")
            return ""
